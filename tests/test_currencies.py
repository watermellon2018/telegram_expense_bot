"""Currency contracts: immutable snapshots, isolation, fallback and concurrent FX lookup."""

import asyncio
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import AsyncMock, MagicMock, patch

import pandas as pd
import pytest

from handlers.currency import currency_keyboard, format_snapshot
from handlers.export import _build_excel_file
from utils import currencies as c
from utils import currency_rates as rates
from utils import currency_reporting as reporting
from utils.parsing import parse_add_command


def test_currency_catalog_and_picker():
    assert set(c.CURRENCIES) == set('RUB USD EUR JPY CNY THB AED TRY GBP KZT GEL KRW'.split())
    markup = currency_keyboard('cur_', ['JPY', 'USD', 'JPY'], 'JPY', other=True)
    assert [row[0].callback_data for row in markup.inline_keyboard] == ['cur_JPY', 'cur_USD', 'cur_other']
    assert 'Японская иена' in markup.inline_keyboard[0][0].text


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '-Infinity', '0', '-1', '1000000000000', 'no'])
def test_invalid_amounts(value):
    with pytest.raises(c.CurrencyError):
        c.parse_amount(value)


def test_parser_default_and_explicit_currency():
    assert parse_add_command('1500 еда обед') == {'amount': Decimal('1500'), 'category': 'еда', 'description': 'обед'}
    explicit = parse_add_command('/add 15,25 usd еда ужин')
    assert explicit == {'amount': Decimal('15.25'), 'currency': 'USD', 'category': 'еда', 'description': 'ужин'}
    assert parse_add_command('15000 KRW еда')['currency'] == 'KRW'
    with pytest.raises(c.CurrencyError):
        parse_add_command('15 XYZ еда')
    assert parse_add_command('15 gym membership')['category'] == 'gym'


@pytest.fixture
def account(monkeypatch):
    monkeypatch.setattr(c, 'require_permission', AsyncMock())
    monkeypatch.setattr(c, 'get_reporting_currency', AsyncMock(return_value='USD'))
    monkeypatch.setattr(c, 'get_input_currency', AsyncMock(return_value='JPY'))


@pytest.mark.asyncio
async def test_identity_never_fetches_rates(account):
    with patch.object(rates, 'get_rate', AsyncMock()) as request:
        money = await c.prepare_money(1, 2, '30.10', 'USD')
    request.assert_not_called()
    assert money['reporting_amount'] == Decimal('30.10')
    assert money['fx_source'] == 'identity'


@pytest.mark.asyncio
async def test_dated_snapshot_and_default_are_preserved(account):
    operation_date = date.today() - timedelta(days=10)
    effective = operation_date - timedelta(days=1)
    with patch.object(rates, 'get_rate', AsyncMock(return_value=(Decimal('0.006'), effective))) as request:
        money = await c.prepare_money(1, 2, '1500', operation_date=operation_date)
    request.assert_awaited_once_with('JPY', 'USD', operation_date)
    assert money['amount'] == Decimal('1500')
    assert money['currency'] == 'JPY'
    assert money['reporting_amount'] == Decimal('9.00')
    assert money['operation_date'] == operation_date
    assert money['fx_date'] == effective
    assert money['fx_source'] == 'cbr'


@pytest.mark.asyncio
async def test_fallback_uses_exact_context_and_pair(account):
    with patch.object(rates, 'get_rate', AsyncMock(side_effect=c.RateUnavailable())), \
         patch.object(c, 'get_fallback_rate', AsyncMock(return_value=Decimal('0.005'))) as fallback:
        money = await c.prepare_money(7, 42, '1500')
    fallback.assert_awaited_once_with(7, 42, 'JPY', 'USD')
    assert money['reporting_amount'] == Decimal('7.50')
    assert money['fx_source'] == 'manual'
    assert 'резервный курс' in format_snapshot(money)


@pytest.mark.asyncio
async def test_no_rate_is_explicit_error(account):
    with patch.object(rates, 'get_rate', AsyncMock(side_effect=c.RateUnavailable())), \
         patch.object(c, 'get_fallback_rate', AsyncMock(return_value=None)):
        with pytest.raises(c.RateUnavailable, match='резервный'):
            await c.prepare_money(1, None, 1500)


@pytest.mark.asyncio
@pytest.mark.parametrize('code', ['JPY', 'KRW'])
async def test_whole_unit_currencies_reject_fraction(account, code):
    with pytest.raises(c.CurrencyError, match='знаков'):
        await c.prepare_money(1, None, '1.5', code)


def test_cbr_nominal_and_effective_date():
    xml = b'''<ValCurs Date="18.09.2026"><Valute><CharCode>JPY</CharCode><Nominal>100</Nominal><Value>60,00</Value></Valute><Valute><CharCode>USD</CharCode><Nominal>1</Nominal><Value>90,00</Value></Valute><Valute><CharCode>KRW</CharCode><Nominal>1000</Nominal><Value>70,00</Value></Valute></ValCurs>'''
    effective, quotes = rates.parse_cbr_xml(xml, date(2026, 9, 19))
    assert effective == date(2026, 9, 18)
    assert quotes['JPY'] == Decimal('0.6')
    assert quotes['KRW'] == Decimal('0.07')
    assert quotes['RUB'] == 1
    with pytest.raises(c.RateUnavailable):
        rates.parse_cbr_xml(xml, date(2026, 9, 17))


@pytest.mark.asyncio
async def test_cross_rate_without_ruble_reporting():
    with patch.object(rates, 'get_daily_rates', AsyncMock(return_value=(date.today(), {'USD': Decimal(90), 'JPY': Decimal('0.6')}))):
        result, _ = await rates.get_rate('USD', 'JPY', date.today())
    assert result == 150


@pytest.mark.asyncio
async def test_concurrent_requests_share_one_fetch(monkeypatch):
    rates._inflight.clear()
    rates._memory.clear()
    gate = asyncio.Event()
    calls = []

    async def fetch(requested):
        calls.append(requested)
        await gate.wait()
        return requested, {'USD': Decimal(90)}

    monkeypatch.setattr(rates, '_load_daily', fetch)
    tasks = [asyncio.create_task(rates.get_daily_rates(date.today())) for _ in range(20)]
    await asyncio.sleep(0)
    await asyncio.sleep(0)
    gate.set()
    results = await asyncio.gather(*tasks)
    assert len(calls) == 1
    assert all(item == results[0] for item in results)


@pytest.mark.asyncio
async def test_context_change_rejects_prepared_money():
    conn = MagicMock(fetchrow=AsyncMock(return_value={'reporting_currency': 'EUR', 'role': 'editor'}))
    with pytest.raises(c.CurrencyError, match='изменилась'):
        await c.validate_money_context(conn, 1, 42, {'reporting_currency': 'USD'})
    assert 'FOR UPDATE' in conn.fetchrow.call_args.args[0]


@pytest.mark.asyncio
async def test_only_owner_can_change_project_currency():
    conn = MagicMock(fetchrow=AsyncMock(return_value={'reporting_currency': 'USD', 'role': 'viewer'}))
    with pytest.raises(PermissionError):
        await c._context(conn, 1, 42, lock=True, owner=True)


@pytest.mark.asyncio
async def test_legacy_summary_scopes_personal_and_income_dates():
    with patch.object(reporting, 'require_permission', AsyncMock()), \
         patch.object(reporting.db, 'fetchrow', AsyncMock(return_value={'count': 2, 'total': Decimal('5000')})) as fetch:
        summary = await reporting.get_legacy_summary(7, None, year=2026, month=9)
    expense_sql, income_sql = [call.args[0] for call in fetch.await_args_list]
    assert 'e.user_id = $1 AND e.project_id IS NULL' in expense_sql
    assert 'e.currency IS NULL' in expense_sql and 'e.deleted_at IS NULL' in expense_sql
    assert 'e.income_date' in income_sql and 'e.date' not in income_sql
    assert summary['expenses']['total'] == 5000
    assert 'отдельно от итогов' in reporting.format_legacy_summary(summary)


def test_export_preserves_originals_and_excludes_legacy_from_total(tmp_path):
    frame = pd.DataFrame([
        {'date': date(2026, 9, 1), 'month': 9, 'category': 'еда', 'description': '', 'original_amount': 1500, 'currency': 'JPY', 'amount': 10.0, 'reporting_currency': 'USD', 'fx_rate': .006666},
        {'date': date(2026, 9, 1), 'month': 9, 'category': 'еда', 'description': '', 'original_amount': Decimal('9999.50'), 'currency': None, 'amount': None, 'reporting_currency': None, 'fx_rate': None},
    ])
    path = tmp_path / 'export.xlsx'
    _build_excel_file(frame, str(path), 9)
    workbook = pd.ExcelFile(path)
    assert 'Без указанной валюты' in workbook.sheet_names
    assert len(pd.read_excel(path, sheet_name='Все расходы')) == 2
    assert pd.read_excel(path, sheet_name='Общая статистика').iloc[0]['Значение'] == 10
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    sheet = workbook['Все расходы']
    column = [cell.value for cell in sheet[1]].index('original_amount') + 1
    assert sheet.cell(3, column).data_type == 'n'
    assert sheet.cell(3, column).value == 9999.5
    workbook.close()


@pytest.mark.asyncio
async def test_legacy_threshold_does_not_compare_untyped_amounts():
    from utils.project_notifier import _should_notify
    settings = {'expense_notify_mode': 'large_only', 'large_expense_threshold': 1000, 'threshold_currency': 'USD'}
    assert not _should_notify(settings, Decimal(1500), 'JPY')
    assert _should_notify(settings, Decimal(1500), 'USD')
    settings['threshold_currency'] = None
    assert not _should_notify(settings, Decimal(1500), 'USD')


@pytest.mark.asyncio
async def test_comparison_reads_only_reporting_amounts_and_checks_access():
    from utils import incomes, permissions

    async def fetch(sql, *args):
        assert args == (42, '7', 2026)
        if 'FROM reporting_incomes' in sql:
            return [{'month': 9, 'total': Decimal('20.00')}]
        assert 'FROM reporting_expenses' in sql
        return [{'month': 9, 'total': Decimal('10.00')}]

    with patch.object(incomes.db, 'fetch', AsyncMock(side_effect=fetch)) as query, \
         patch.object(permissions, 'require_permission', AsyncMock()) as permission:
        data = await incomes.get_yearly_income_vs_expense(7, 2026, 42)
        assert data == {'incomes': {9: 20.0}, 'expenses': {9: 10.0}}
        permission.assert_awaited_once_with(7, 42, permissions.Permission.VIEW_STATS)
        query.reset_mock()
        permission.side_effect = PermissionError
        with pytest.raises(PermissionError):
            await incomes.get_yearly_income_vs_expense(7, 2026, 42)
        query.assert_not_awaited()


@pytest.mark.asyncio
async def test_writer_rejected_if_role_became_viewer():
    conn = MagicMock(fetchrow=AsyncMock(return_value={'reporting_currency': 'USD', 'role': 'viewer'}))
    with pytest.raises(PermissionError):
        await c.validate_money_context(conn, 1, 42, {'reporting_currency': 'USD'}, writer=True)


def test_suggestion_keys_bind_currency_and_project():
    from utils import pattern_detector as patterns
    base = {'project_id': 42, 'currency': 'JPY', 'category_id': 5, 'amount': 100,
            'comment': 'кофе', 'frequency_type': 'daily'}
    other_currency = {**base, 'currency': 'KRW'}
    other_project = {**base, 'project_id': None}
    assert len({patterns.pattern_hash(p) for p in (base, other_currency, other_project)}) == 3
    cache = {}
    for pattern in (base, other_currency, other_project):
        patterns.save_pattern_to_cache(cache, '7', pattern)
    assert len(cache['rec_patterns']) == 3


def test_full_pdf_preserves_expense_rounded_to_zero():
    from io import BytesIO

    import matplotlib.pyplot as plt

    from utils import report_generator
    from utils.formatting import format_month_expenses

    day = date(2026, 9, 19)
    frame = pd.DataFrame([{'date': day, 'amount': 0.0, 'category': 'food', 'description': ''}])
    output = BytesIO()
    try:
        report_generator._render_full_report(frame, frame.iloc[:0], day, output, currency='USD')
        assert output.getvalue().startswith(b'%PDF')
        assert len(output.getvalue()) > 10000
        text = format_month_expenses({'count': 1, 'total': 0, 'by_category': {'food': 0}}, 9, 2026)
        assert 'Количество транзакций: 1' in text
    finally:
        plt.close('all')
